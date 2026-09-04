from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

from src.config import BASE_DIR, DB_PATH
from src.screener.engine import ScreenerEngine
from src.analytics.peer import compute_peer_percentiles, load_peer_groups, PEER_METRICS

GREEN = PatternFill("solid", fgColor="D9EAD3")
RED = PatternFill("solid", fgColor="F4CCCC")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
GOLD = PatternFill("solid", fgColor="FFD966")


def export_screener_workbook(path: str | Path = BASE_DIR / "output" / "screener_output.xlsx", db_path: str | Path = DB_PATH) -> Path:
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    engine = ScreenerEngine(db_path)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name in engine.config["presets"]:
            frame = engine.run_preset_screener(name)
            preferred = ["company_id", "ticker", "company_name", "broad_sector", "year"]
            metrics = [c for c in frame.columns if c not in preferred and c != "composite_quality_score"][:20]
            frame[preferred + metrics + ["composite_quality_score"]].to_excel(writer, sheet_name=name[:31], index=False)
    workbook = pd.ExcelFile(path)
    from openpyxl import load_workbook
    book = load_workbook(path)
    for sheet in book.worksheets:
        for cell in sheet[1]: cell.font = Font(bold=True)
        if sheet.max_row < 2:
            continue
        end = get_column_letter(sheet.max_column)
        sheet.conditional_formatting.add(f"A2:{end}{sheet.max_row}", CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=GREEN))
    book.save(path)
    return path


def export_peer_workbook(
    path: str | Path = BASE_DIR / "output" / "peer_comparison.xlsx",
    db_path: str | Path = DB_PATH, workbook_path: str | Path | None = None,
) -> Path:
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    percentiles = compute_peer_percentiles(workbook_path, db_path)
    groups = load_peer_groups(workbook_path, db_path)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for group_name, assignments in groups.groupby("peer_group_name"):
            ids = assignments.company_id.tolist()
            rows = percentiles[percentiles.company_id.isin(ids)].pivot_table(
                index="company_id", columns="metric", values=["value", "percentile_rank"]
            ).reset_index()
            rows.columns = ["company_id"] + [f"{a}_{b}" for a, b in rows.columns.tolist()[1:]]
            names = assignments.merge(pd.read_sql_query("SELECT company_id, company_name FROM companies", __import__("sqlite3").connect(db_path)), on="company_id")
            rows = names[["company_id", "company_name"]].merge(rows, on="company_id")
            median = rows.drop(columns=["company_id"], errors="ignore").select_dtypes("number").median().to_frame().T
            median.insert(0, "company_name", "Peer group median"); median.insert(0, "company_id", "MEDIAN")
            pd.concat([rows, median], ignore_index=True).to_excel(writer, sheet_name=str(group_name)[:31], index=False)
    from openpyxl import load_workbook
    book = load_workbook(path)
    for sheet in book.worksheets:
        for cell in sheet[1]: cell.font = Font(bold=True)
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            if row[0].value == "MEDIAN":
                for cell in row: cell.fill = GOLD
            for cell in row:
                if isinstance(cell.value, float):
                    cell.fill = GREEN if cell.value >= .75 else RED if cell.value <= .25 else YELLOW
    book.save(path)
    return path